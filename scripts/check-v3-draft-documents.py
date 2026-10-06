"""Audit continuation DOCX content. This does not certify visual quality."""
import argparse
import json
import re
from pathlib import Path
from docx import Document

parser = argparse.ArgumentParser()
parser.add_argument('--documents', type=Path, default=Path('work/v3-docs'))
args = parser.parse_args()
root = Path(__file__).resolve().parents[1]


def norm(text):
    return re.sub(r'[^a-z0-9]', '', str(text).lower())


count = 0
for path in sorted(args.documents.glob('ai-*-week-*-*.docx')):
    stem, role = path.stem.rsplit('-', 1)
    source = json.loads((root / 'curriculum-v3' / f'{stem}.json').read_text())
    doc = Document(path)
    text = norm(' '.join(p.text for p in doc.paragraphs) + ' ' + ' '.join(c.text for t in doc.tables for r in t.rows for c in r.cells))
    expected = [source['title']]
    if role == 'lesson':
        expected += source['goals'] + [source['prerequisite']['prompt']]
        for section in source['lesson']:
            expected.append(section['title'])
            for block in section['blocks']:
                kind = block['type']
                if kind in ('paragraph', 'code'):
                    expected.append(block['text'])
                elif kind == 'table':
                    expected += block['headers'] + [v for row in block['rows'] for v in row]
                elif kind == 'diagram':
                    expected += [block['title'], block['caption']]
                elif kind == 'response':
                    expected.append(block['label'])
        expected += [v for row in source['vocabulary'] for v in row]
    elif role == 'workbook':
        expected += [v for t in source['workbook'] for v in (t['title'], t['prompt'])]
    else:
        assert role == 'guide'
        teacher = source['teacher']
        expected += source['goals'] + [source['prerequisite']['prompt'], source['prerequisite']['answer']]
        expected += [teacher[k] for k in ('background', 'guidedAnswer', 'assessment', 'support', 'next')]
        expected += teacher['materials'] + [v for row in teacher['sessions'] + teacher['misconceptions'] for v in row]
        expected += [t['answer'] for t in source['workbook']]
    for value in expected:
        assert norm(value) in text, (path.name, str(value)[:100])
    if role != 'guide':
        for task in source['workbook']:
            assert norm(task['answer']) not in text, (path.name, 'Teacher answer found in student document')
    count += 1

assert count, 'No draft documents found'
print(f'{count} DOCX files match canonical content; student files exclude teacher answer explanations. Visual quality is recorded separately in print-review.json.')
