"""Audit DOCX/PDF text and page bounds; render visual-review images.

This creates a pending receipt. It never claims that images were inspected.
"""
import argparse
import hashlib
import json
from pathlib import Path
import re
import shutil
import fitz
from docx import Document

parser=argparse.ArgumentParser()
parser.add_argument('--render',action='store_true')
args=parser.parse_args()
ROOT=Path(__file__).resolve().parents[1]
review={'status':'PENDING visual inspection and editorial review','chapters':{},'pages':0,'boundsIssues':[]}


def norm(text):return re.sub(r'[^a-z0-9]','',str(text).lower())
def sha(path):return hashlib.sha256(path.read_bytes()).hexdigest()


for docx in sorted((ROOT/'work/v3-docs').glob('*.docx')):
    stem,role=docx.stem.rsplit('-',1)
    pdf=ROOT/f'work/v3-pdf/{docx.stem}.pdf'
    assert pdf.is_file(),f'Missing PDF: {pdf}'
    source=ROOT/f'curriculum-v3/{stem}.json'
    w=json.loads(source.read_text())
    doc=Document(docx)
    with fitz.open(pdf) as rendered:
        text=norm(' '.join(page.get_text() for page in rendered))
        expected=[p.text for p in doc.paragraphs if p.text.strip()]
        expected += [c.text for table in doc.tables for row in table.rows for c in row.cells if c.text.strip()]
        for value in expected:
            assert norm(value) in text,(pdf.name,'DOCX/PDF text mismatch',value[:120])
        if role!='guide':
            for task in w['workbook']:
                assert norm(task['answer']) not in text,(pdf.name,'teacher explanation in student PDF')
        folder=ROOT/f'work/v3-render/{docx.stem}'
        if args.render:
            folder.mkdir(parents=True,exist_ok=True)
            shutil.copyfile(pdf,folder/pdf.name)
        pages=[]
        for number,page in enumerate(rendered,1):
            for word in page.get_text('words'):
                x0,y0,x1,y1,*_=word
                if x0 < -1 or y0 < -1 or x1 > page.rect.width+1 or y1 > page.rect.height+1:
                    review['boundsIssues'].append([pdf.name,number,word[:5]])
            if args.render:
                path=folder/f'page-{number}.png'
                page.get_pixmap(matrix=fitz.Matrix(1.3,1.3),alpha=False).save(path)
                pages.append(sha(path))
            review['pages']+=1
        id=f'{w["level"]}/{w["week"]}'
        chapter=review['chapters'].setdefault(id,{'sourceHash':sha(source),'visualReview':'PENDING','documents':{}})
        chapter['documents'][role]={'docx':sha(docx),'pdf':sha(pdf),'pageCount':len(rendered),'pages':pages}

assert len(review['chapters'])==112
assert all(set(r['documents'])=={'lesson','workbook','guide'} for r in review['chapters'].values())
(ROOT/'work/v3-draft-pdf-audit.json').write_text(json.dumps(review,indent=2)+'\n')
print(f'336 PDFs match their DOCX content across {review["pages"]} pages; {len(review["boundsIssues"])} text-bound issues. Visual/editorial review remains pending.')
