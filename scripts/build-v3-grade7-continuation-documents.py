"""Generate matched teaching documents from the canonical Curriculum 3 source.

Run with the bundled Python runtime. Render every output with render_docx.py
and inspect every page before adding it to the protected release manifest.
"""
import argparse,json
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

parser=argparse.ArgumentParser();parser.add_argument('--out',required=True);parser.add_argument('--diagrams',required=True);args=parser.parse_args()
root=Path(__file__).resolve().parents[1];out=Path(args.out).resolve();out.mkdir(parents=True,exist_ok=True);diagrams=Path(args.diagrams).resolve()

def new_doc(w,role):
 level=json.loads((root/'curriculum-v3/progression.json').read_text())['levels'][int(w['level'].split('-')[1])-1]
 code=w['level'].upper().replace('-', ' ')
 d=Document();sec=d.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
 sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.75)
 for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
  st=d.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
  for border in st.element.xpath('.//w:pBdr'):border.getparent().remove(border)
  st.element.get_or_add_rPr().rFonts.set(qn('w:ascii'),'Arial');st.element.rPr.rFonts.set(qn('w:hAnsi'),'Arial')
 d.styles['Normal'].font.size=Pt(11);d.styles['Normal'].paragraph_format.line_spacing=1.13;d.styles['Normal'].paragraph_format.space_after=Pt(7)
 for name,size in [('Title',22),('Heading 1',17),('Heading 2',13),('Heading 3',11)]:
  d.styles[name].font.size=Pt(size);d.styles[name].paragraph_format.keep_with_next=True
 footer=sec.footer.paragraphs[0];footer.alignment=1
 footer.add_run(f'AI Academy   {code}   Week {w["week"]}   {role.title()}   ').font.size=Pt(9)
 field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
 d.add_heading(w['title'],0);d.add_paragraph(f'{code} {level["name"]}   Grade {level["grade"]}   Week {w["week"]}   {role.title()}',style='Subtitle')
 return d

def paragraph(d,text):d.add_paragraph(text)
def table(d,headers,rows):
 t=d.add_table(rows=1,cols=len(headers));t.style='Table Grid';t.autofit=False
 for c,label in zip(t.rows[0].cells,headers):
  c.text=label
  shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E9EFF5');c._tc.get_or_add_tcPr().append(shade)
  for r in c.paragraphs[0].runs:r.bold=True
 hdr=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(hdr)
 for row in rows:
  for c,value in zip(t.add_row().cells,row):c.text=str(value)
 for row in t.rows:
  pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
  for c in row.cells:
   for p in c.paragraphs:
    p.paragraph_format.space_after=Pt(5);p.paragraph_format.space_before=Pt(5)
    for r in p.runs:r.font.size=Pt(10)
 d.add_paragraph()

def response(d,label,lines):
 d.add_paragraph(label)
 for _ in range(lines):
  p=d.add_paragraph(' ');p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1
  p.paragraph_format.line_spacing=Pt(15)
  border=OxmlElement('w:pBdr');bottom=OxmlElement('w:bottom');bottom.set(qn('w:val'),'single');bottom.set(qn('w:sz'),'3');bottom.set(qn('w:color'),'B9C1CA');border.append(bottom);between=OxmlElement('w:between');between.attrib.update(bottom.attrib);border.append(between);p._p.get_or_add_pPr().append(border)

def lesson_doc(w):
 d=new_doc(w,'student lesson');paragraph(d,'This chapter explains how the weekly idea works. Read the examples, complete the supported explanation, then use the workbook to show that you can apply the idea to a new case.')
 d.add_heading('What you will learn',1)
 for goal in w['goals']:d.add_paragraph(goal,style='List Bullet')
 d.add_heading('Recall before reading',2);paragraph(d,w['prerequisite']['prompt'])
 diagram=0
 for i,s in enumerate(w['lesson']):
  if i in (1,2):d.add_page_break()
  d.add_heading(s['title'],1)
  for b in s['blocks']:
   if b['type']=='paragraph':paragraph(d,b['text'])
   elif b['type']=='code':
    for index,line in enumerate(b['text'].splitlines()):
     p=d.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.05;p.paragraph_format.keep_with_next=index<len(b['text'].splitlines())-1
     run=p.add_run(line or ' ');run.font.name='Courier New';run.font.size=Pt(10)
    d.add_paragraph()
   elif b['type']=='table':table(d,b['headers'],b['rows'])
   elif b['type']=='diagram':
    diagram+=1;d.add_heading(b['title'],2)
    p=d.add_paragraph();p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(diagrams/f'{w["level"]}-week-{w["week"]:02}-diagram-{diagram}.png'),width=Inches(6.6))
    paragraph(d,b['caption'])
   elif b['type']=='response':response(d,b['label'],b['lines'])
 d.add_heading('Words to use accurately',1);table(d,['Word','Meaning'],w['vocabulary'])
 return d

def workbook_doc(w):
 d=new_doc(w,'workbook');paragraph(d,'Write your name and date. Complete the supported tasks first. Attempt the new case independently and record any help you use. Your teacher has the answer explanations.')
 for i,t in enumerate(w['workbook']):
  if i and i%2==0:d.add_page_break()
  d.add_heading(t['id']+' '+t['title'],1);paragraph(d,t['prompt']);response(d,'My reasoning and evidence',t['lines'])
 d.add_paragraph('Support used   None   Hint   Teacher explanation')
 return d

def teacher_doc(w):
 d=new_doc(w,'teacher guide');t=w['teacher'];paragraph(d,t['background'])
 d.add_heading('Learning outcomes',1)
 for goal in w['goals']:d.add_paragraph(goal,style='List Bullet')
 d.add_heading('Preparation',1)
 for item in t['materials']:d.add_paragraph(item,style='List Bullet')
 d.add_heading('Prerequisite check and answer',1);paragraph(d,w['prerequisite']['prompt']);paragraph(d,w['prerequisite']['answer'])
 d.add_page_break();d.add_heading('Teaching sequence',1)
 for title,body in t['sessions']:d.add_heading(title,2);paragraph(d,body)
 d.add_page_break();d.add_heading('Solutions to student chapter questions',1);paragraph(d,t['guidedAnswer'])
 d.add_heading('Workbook answer explanations',1)
 for task in w['workbook']:d.add_heading(task['id']+' '+task['title'],2);paragraph(d,task['answer'])
 d.add_page_break();d.add_heading('Respond to misconceptions',1)
 for claim,repair in t['misconceptions']:d.add_heading(claim,2);paragraph(d,repair)
 for title,key in [('Assess independent understanding','assessment'),('Support and extension','support'),('Connect to the next week','next')]:d.add_heading(title,1);paragraph(d,t[key])
 return d

for file in sorted((root/'curriculum-v3').glob('ai-2-week-*.json')):
 w=json.loads(file.read_text())
 if w['week'] < 10:continue
 for role,build in [('lesson',lesson_doc),('workbook',workbook_doc),('guide',teacher_doc)]:
  d=build(w);dest=out/f'{w["level"]}-week-{w["week"]:02}-{role}.docx';d.save(dest);print(dest)
