"""Export matching editable Word editions; requires python-docx and Pillow."""
from reading_style import segments
from pathlib import Path
import json,os
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from PIL import Image
ROOT=Path(__file__).resolve().parents[1]
D=json.loads((ROOT/'content/programme.json').read_text())
def stem(w):return ('bridge-' if w['id'].startswith('python') else f"ai-{w['level']}-")+f"week-{w['week']:02}"
def add_table(doc,rows):
 if not rows:return
 count=max(map(len,rows));t=doc.add_table(rows=0,cols=count);t.style='Light Shading Accent 1'
 for row in rows:
  cells=t.add_row().cells
  for i,v in enumerate(row):cells[i].text=str(v)
 t.rows[0]._tr.get_or_add_trPr().append(__import__('docx').oxml.OxmlElement('w:tblHeader'))
 return t
def block(doc,b):
 f=b.get('presentation')
 if f:
  if f['kind']=='section-label':return
  if f['kind']=='checklist':
   for item in f['items']:doc.add_paragraph(item,style='List Bullet')
   doc.add_paragraph(f['tail']);return
  if f.get('label'):doc.add_heading(f['label'],3)
  if f.get('context'):doc.add_paragraph(f['context'])
  colors={'concept':'EDF7F3','example':'F1F6FC','question':'F5F1FB','caution':'FFF5E6','note':'F1F6F8','definition':'F4F9FC','step':'F4F9FC'}
  for text in f.get('paragraphs',[b['text']]):
   p=doc.add_paragraph();p.paragraph_format.line_spacing=1.4;p.paragraph_format.space_after=Pt(9)
   if f['kind'] in colors:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),colors[f['kind']]);p._p.get_or_add_pPr().append(shade)
   for part,strong in segments(text,f.get('emphasis',[])):p.add_run(part).bold=strong or f['kind']=='eyebrow'
  return
 kind=b['type'];txt=b.get('text','')
 if kind=='gallery':
  for cell in b['cells']:
   for child in cell:block(doc,child)
 elif kind=='image':
  p=ROOT/'assets'/(b['src'].split('/')[-1]+'.png')
  with Image.open(p) as im:width,height=im.size
  size=min(6,width/110,(6.5 if b.get("academy_visual") else 3.5)*width/height)
  shape=doc.add_picture(str(p),width=Inches(size));shape._inline.docPr.set('descr',b.get('alt','Concept illustration'))
 elif kind=='diagram':
  doc.add_heading(b.get('title','Trace the evidence'),3)
  if b.get('rows'):
   table=add_table(doc,[b.get('columns', ['Stage','Evidence','Result'])]+b['rows'])
   if b.get('columns')==['Step','Evidence and reasoning']:
    table.autofit=False
    for row in table.rows:
     row.cells[0].width=Inches(.55);row.cells[1].width=Inches(6.55)
  if b.get('values'):
   doc.add_paragraph(b.get('key',''));add_table(doc,b['values'])
  if b.get('series'):add_table(doc,[['Series']+list(map(str,b['xValues']))]+[[s['label']]+list(map(str,s['values'])) for s in b['series']])
  doc.add_paragraph(b.get('caption',''))
 elif kind=='table':add_table(doc,b['rows'])
 elif kind=='code':
  p=doc.add_paragraph();r=p.add_run(txt);r.font.name='Courier New';r.font.size=Pt(9)
 elif kind=='response':
  doc.add_paragraph(b.get('label','My reasoning and evidence'),style='Heading 3')
  for _ in range(min(b.get('lines',5),8)):doc.add_paragraph('________________________________________________________________')
 elif kind=='sensor':
  doc.add_heading('The sensor measures. The rule decides.',3)
  doc.add_paragraph('This model uses a made-up light scale: 0 means darker and 100 brighter. Rule: if reading < 30, lamp ON; otherwise lamp OFF. Trace readings 18, 29 and 30. Explain what measured the light and what chose the action. The HTML lesson includes the interactive slider.')
 elif kind in ['heading','title','subheading']:doc.add_heading(txt,3)
 else:doc.add_paragraph(txt)
manifest=[]
for w in D['weeks']+D['python_bridge']:
 if os.environ.get('CHECK_WEEKS') and w['id'] not in os.environ['CHECK_WEEKS'].split(','):continue
 for role,key in [('student','sections'),('teacher','teacher'),('workbook','workbook')]:
  doc=Document();sec=doc.sections[0];sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.7)
  normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(11);normal.paragraph_format.space_after=Pt(7)
  for n in ['Heading 1','Heading 2','Heading 3']:doc.styles[n].font.color.rgb=RGBColor.from_string('174C69')
  doc.add_paragraph(f"AI ACADEMY · {role.upper()} · {w['id']}")
  doc.add_heading(w['title'],0);doc.add_paragraph('Project: '+w['project'])
  if role=='student':doc.add_heading('Before you start',1);doc.add_paragraph(w['recall'])
  for i,s in enumerate(w[key],1):
   doc.add_heading(f"{i:02}. {s['title']}",1)
   if s.get('reading_intro'):doc.add_paragraph(s['reading_intro'])
   for b in s['blocks']:block(doc,b)
  path=ROOT/'downloads'/role/(stem(w)+'.docx');path.parent.mkdir(parents=True,exist_ok=True);doc.save(path)
  manifest.append({'week':w['id'],'role':role,'file':str(path.relative_to(ROOT))})
 print('Word exports:',len(manifest),flush=True) if len(manifest)% 90==0 else None
if not os.environ.get('CHECK_WEEKS'):(ROOT/'reports/word-manifest.json').write_text(json.dumps(manifest,indent=2))
print('Exported',len(manifest),'Word documents')
